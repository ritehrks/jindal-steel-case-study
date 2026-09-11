# StainlessCarbon Round 1 deck - NATIVE PowerPoint elements (fully editable)
# Light consulting style matching the approved HTML design.
# All numbers trace to research/22-canonical-model-v1.md.

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

NAVY  = RGBColor(0x12,0x3F,0x63); BLUE  = RGBColor(0x1B,0x6D,0xB0)
TEXT  = RGBColor(0x1F,0x29,0x37); GRAY  = RGBColor(0x5B,0x65,0x70)
FAINT = RGBColor(0x8A,0x94,0xA0); LINE  = RGBColor(0xC9,0xD2,0xDC)
LINE2 = RGBColor(0xE2,0xE8,0xEF); THBG  = RGBColor(0xEE,0xF2,0xF6)
BOX   = RGBColor(0xF5,0xF7,0xFA); WHITE = RGBColor(0xFF,0xFF,0xFF)
RED   = RGBColor(0xB3,0x39,0x2F); GREEN = RGBColor(0x1D,0x7A,0x4F)
AMBERD= RGBColor(0xA8,0x72,0x0E); STEEL = RGBColor(0x7B,0x87,0x94)
GBOX  = RGBColor(0xF2,0xF7,0xEE); GLINE = RGBColor(0xCF,0xE0,0xC5)

URL_TOOL = 'https://ritehrks.github.io/jindal-steel-case-study/website/'
URL_APX  = 'https://ritehrks.github.io/jindal-steel-case-study/appendix.pdf'
HERE = os.path.dirname(os.path.abspath(__file__))

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
M = Inches(0.32)

def bg(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE

def box(s, x, y, w, h, fill=BOX, line=None, lw=0.75):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp

def txt(s, x, y, w, h, paras, size=10, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space_after=2, line_spacing=1.0, font='Segoe UI'):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(1.5); tf.margin_top = tf.margin_bottom = Pt(0.5)
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        if isinstance(para, str): para = [(para, {})]
        for t, st in para:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color); f.name = st.get('font', font)
            if st.get('italic'): f.italic = True
            if st.get('link'): r.hyperlink.address = st['link']
    return tb

def hline(s, x, y, w, color=LINE2, weight=0.8):
    return box(s, x, y, w, Pt(weight), fill=color)

def sec(s, x, y, w, label):
    txt(s, x, y, w, Inches(0.2),
        [[(label.upper(), {'size': 9, 'bold': True, 'color': BLUE})]], space_after=0)
    hline(s, x, y+Inches(0.21), w)
    return y + Inches(0.30)

def set_cell(cell, paras, size=8.5, color=TEXT, bold=False, align=PP_ALIGN.LEFT, fill=WHITE):
    cell.fill.solid(); cell.fill.fore_color.rgb = fill
    if paras == '' or paras == []:
        paras = [[(' ', {'size': 4})]]  # keep empty cells from inflating row height
    tf = cell.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4); tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(para, str): para = [(para, {})]
        for t, st in para:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color); f.name = 'Segoe UI'

def mktable(s, rows, cols, x, y, w, h):
    tb = s.shapes.add_table(rows, cols, x, y, w, h).table
    el = tb._tbl.find(qn('a:tableStyleId'))
    if el is not None: tb._tbl.remove(el)
    return tb

def footer(s, pageno, src_text):
    hline(s, M, Inches(6.98), prs.slide_width-2*M, color=LINE2)
    txt(s, M, Inches(7.04), Inches(12.2), Inches(0.42),
        [[(src_text+'  ', {'size': 7, 'color': FAINT}),
          ('Live prototype: ', {'size': 7, 'bold': True, 'color': GRAY}),
          ('ritehrks.github.io/jindal-steel-case-study/website', {'size': 7, 'color': BLUE, 'link': URL_TOOL}),
          ('   ·   ', {'size': 7, 'color': FAINT}),
          ('Research appendix: ', {'size': 7, 'bold': True, 'color': GRAY}),
          ('ritehrks.github.io/jindal-steel-case-study/appendix.pdf', {'size': 7, 'color': BLUE, 'link': URL_APX})]],
        line_spacing=1.15)
    txt(s, Inches(12.85), Inches(7.04), Inches(0.35), Inches(0.22),
        [[(str(pageno), {'size': 8, 'color': GRAY})]], align=PP_ALIGN.RIGHT)

# =====================================================================
# SLIDE 1
# =====================================================================
s = prs.slides.add_slide(blank); bg(s)

txt(s, M, Inches(0.10), Inches(10.0), Inches(0.66),
    [[("Stainless steel's carbon is an ", {'size': 14.5, 'color': NAVY}),
      ('alloy-chain problem', {'size': 14.5, 'bold': True, 'color': NAVY}),
      (' — ferrochrome and nickel drive 60–80% of the footprint,', {'size': 14.5, 'color': NAVY})],
     [('and no existing calculator models them. ', {'size': 14.5, 'color': NAVY}),
      ('We built and deployed one that does.', {'size': 14.5, 'bold': True, 'color': NAVY})]],
    space_after=1, line_spacing=1.08)
txt(s, Inches(10.35), Inches(0.16), Inches(2.68), Inches(0.6),
    [[('Jindal Stainless · "Spark the Rising Curve"', {'size': 7.5, 'color': FAINT})],
     [('Problem Statement 3 · Round 1 · 11 Sep 2026', {'size': 7.5, 'color': FAINT})]],
    align=PP_ALIGN.RIGHT, space_after=1.5)
box(s, M, Inches(0.80), prs.slide_width-2*M, Pt(2.2), fill=NAVY)

colY = Inches(0.94)
# ---- LEFT ----
LX, LW = M, Inches(2.95)
y = sec(s, LX, colY, LW, 'Why existing tools fall short')
gaps = [
    ('They ignore the alloys. ', 'Odisha ferrochrome runs 5.4–5.9 tCO₂/t (2.5× the global average); NPI carries 60–85 tCO₂/t of contained nickel. Carbon-steel tools model neither.'),
    ('They assume Western grids. ', "Typical tools use 0.2–0.4 t/MWh; India's grid is 0.710 and captive subcritical coal 1.045 t/MWh (CEA v21.0)."),
    ('They allow infeasible answers. ', "Sliding scrap to 100% ignores Cu/Sn tramp-element limits and India's 23–25% scrap availability; the practical ceiling is ~80%."),
    ('They ignore the money. ', 'From 2026 the EU prices every imported tonne (CBAM). Default vs verified data is a €370–400/t swing by 2030.'),
]
for head, body in gaps:
    txt(s, LX, y, LW, Inches(0.66),
        [[(head, {'size': 8.8, 'bold': True, 'color': NAVY}), (body, {'size': 8.3, 'color': GRAY})]],
        line_spacing=1.02)
    y += Inches(0.68)
tb = mktable(s, 6, 4, LX, y+Inches(0.04), LW, Inches(1.55))
tb.columns[0].width = Inches(1.28); tb.columns[1].width = Inches(0.70)
tb.columns[2].width = Inches(0.50); tb.columns[3].width = Inches(0.47)
hdr = ['', 'Primetals', 'SSAB', 'Ours']
rows = [('Alloy-chain emissions','–','–','✓'), ('India grid / captive coal','–','–','✓'),
        ('Feasibility guardrails','–','–','✓'), ('CBAM cost per tonne','partial','–','✓'),
        ('Constrained optimizer','partial','–','✓')]
for j,h in enumerate(hdr):
    set_cell(tb.cell(0,j), [[(h, {'size':7,'bold':True,'color':GRAY})]], fill=THBG,
             align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
for i,(a,b,c,d) in enumerate(rows, start=1):
    set_cell(tb.cell(i,0), [[(a, {'size':7.6,'color':TEXT})]], fill=WHITE)
    for j,v in enumerate([b,c,d], start=1):
        col = GREEN if v=='✓' else (AMBERD if v=='partial' else FAINT)
        set_cell(tb.cell(i,j), [[(v, {'size':7.6,'bold':v!='–','color':col})]], align=PP_ALIGN.CENTER, fill=WHITE)
for r in tb.rows: r.height = Inches(0.258)

# ---- CENTER ----
CX, CW = M+LW+Inches(0.22), Inches(5.62)
y = sec(s, CX, colY, CW, 'The tool — working prototype, deployed (link below)')
pic = s.shapes.add_picture(os.path.join(HERE, 'toolshot.png'), CX, y, width=CW)
pic.line.color.rgb = LINE; pic.line.width = Pt(1)
capy = y + Emu(int(CW*0.7)) + Inches(0.06)
txt(s, CX, capy, CW, Inches(0.62),
    [[('Actual product screenshot. ', {'size': 8.2, 'bold': True, 'color': NAVY}),
      ('Four operating levers (scrap, nickel source, renewables, clean fuel) with grade selection; live scope split; scenario plotted against Jindal’s audited FY22–26 trajectory; sensitivity ranking of the next best move. A constrained optimizer and a per-tonne CBAM statement sit below the fold.', {'size': 8.2, 'color': GRAY})]],
    line_spacing=1.05)

# ---- RIGHT ----
RX = CX+CW+Inches(0.22); RW = prs.slide_width - RX - M
y = sec(s, RX, colY, RW, 'How the model works')
box(s, RX, y, RW, Inches(0.66), fill=BOX, line=LINE2)
txt(s, RX+Inches(0.06), y+Inches(0.04), RW-Inches(0.12), Inches(0.6),
    [[('CO₂/t = S1 (fuels, AOD, captive power)', {'size': 8, 'color': TEXT, 'font': 'Consolas'})],
     [('  + S2 (grid electricity × CEA factor)', {'size': 8, 'color': TEXT, 'font': 'Consolas'})],
     [('  + S3 ((1−scrap%) × Σ alloy mass × EF)', {'size': 8, 'color': TEXT, 'font': 'Consolas'})]],
    space_after=1, line_spacing=1.05)
y += Inches(0.76)
tb = mktable(s, 6, 3, RX, y, RW, Inches(1.42))
tb.columns[0].width = Inches(1.35); tb.columns[1].width = Inches(1.05); tb.columns[2].width = RW - Inches(2.40)
facts = [('Key factor','Value','Source'), ('India grid','0.710 t/MWh','CEA v21.0'),
         ('Captive coal power','1.045 t/MWh','CEA heat rate'), ('Odisha ferrochrome','5.4–5.9 t/t','ICDA LCA'),
         ('NPI (contained Ni)','60–85 t/t','ISSF'), ('Grade chemistry','304/316/430/2205','mass balance')]
for i,row in enumerate(facts):
    for j,v in enumerate(row):
        set_cell(tb.cell(i,j), [[(v, {'size':7.4,'bold':i==0,'color':GRAY if i==0 else TEXT})]],
                 fill=THBG if i==0 else WHITE,
                 align=PP_ALIGN.RIGHT if (j==1 and i>0) else PP_ALIGN.LEFT)
for r in tb.rows: r.height = Inches(0.235)
y += Inches(1.54)
y = sec(s, RX, y, RW, 'Practicality built in')
prac = ['Scrap capped at 80% (Cu/Sn tramp-element metallurgy; India scrap supply)',
        'Captive ferrochrome kept inside Scope 1+2 — no double counting with Scope 3',
        'Dual output per ISSF method: recycled-content and end-of-life credit']
for pt_ in prac:
    txt(s, RX+Inches(0.02), y, RW-Inches(0.04), Inches(0.34),
        [[('▪ ', {'size': 7.5, 'color': BLUE}), (pt_, {'size': 8, 'color': TEXT})]], line_spacing=1.0)
    y += Inches(0.335)
box(s, RX, y+Inches(0.03), RW, Inches(0.82), fill=GBOX, line=GLINE)
txt(s, RX+Inches(0.07), y+Inches(0.08), RW-Inches(0.14), Inches(0.74),
    [[('Calibration check: ', {'size': 8.2, 'bold': True, 'color': GREEN}),
      ('with Jindal’s disclosed FY26 inputs (70% scrap, 47% RE), the model returns 1.73 / 2.99 tCO₂e/t against the audited 1.76 / 3.03 — within ~2%.', {'size': 8.2, 'color': RGBColor(0x2C,0x4A,0x22)})]],
    line_spacing=1.05)

# ---- process strip ----
py = Inches(6.06)
hline(s, M, py-Inches(0.06), prs.slide_width-2*M)
steps = [('Charge make-up','scrap + FeCr + Ni units (Scope 3)'),
         ('EAF melt','400–600 kWh/t (Scope 2)'),
         ('AOD refine','process CO₂ (Scope 1)'),
         ('Cast + reheat','FO/LSHS → NG, Bio-LDO'),
         ('Roll + finish','anneal, pickle; propane → H₂'),
         ('Mill gate','1.76 S1+2 · 3.03 total (FY26)'),
         ('EU border','CBAM 2030: €446 default vs €48–74 verified')]
n = len(steps); gapw = Inches(0.14)
sw = Emu(int((prs.slide_width - 2*M - gapw*(n-1)) / n))
x = M
for i,(t1,t2) in enumerate(steps):
    box(s, x, py, sw, Inches(0.62), fill=BOX, line=LINE)
    txt(s, x+Inches(0.05), py+Inches(0.03), sw-Inches(0.1), Inches(0.56),
        [[(t1, {'size': 8, 'bold': True, 'color': NAVY})],
         [(t2, {'size': 7, 'color': GRAY})]], space_after=0.5, line_spacing=0.95)
    if i < n-1:
        txt(s, x+sw-Inches(0.015), py+Inches(0.17), gapw+Inches(0.03), Inches(0.25),
            [[('›', {'size': 10, 'color': FAINT})]], align=PP_ALIGN.CENTER)
    x += sw + gapw

footer(s, 1, 'Source: JSL ESG Factsheets, BRSR and CDP disclosures FY22–26; CEA CO₂ Baseline Database v21.0; BEE PAT; IPCC 2006; worldstainless/ISSF; ICDA ferrochrome LCA; EU Regulations 2023/956 and 2025/2620-21; team analysis (19 research files, 40+ primary sources).')

# =====================================================================
# SLIDE 2
# =====================================================================
s = prs.slides.add_slide(blank); bg(s)
txt(s, M, Inches(0.10), Inches(10.35), Inches(0.66),
    [[("Calibrated to Jindal's audited disclosures, the model maps a practical path from ", {'size': 14.5, 'color': NAVY}),
      ('1.76 to 1.35–1.45 tCO₂e/t', {'size': 14.5, 'bold': True, 'color': NAVY}),
      (' —', {'size': 14.5, 'color': NAVY})],
     [('and verified data protects ', {'size': 14.5, 'color': NAVY}),
      ('₹200–290 crore a year', {'size': 14.5, 'bold': True, 'color': NAVY}),
      (' of CBAM exposure by 2030', {'size': 14.5, 'color': NAVY})]],
    space_after=1, line_spacing=1.08)
txt(s, Inches(10.55), Inches(0.16), Inches(2.48), Inches(0.6),
    [[('Round 1 · Executive summary', {'size': 7.5, 'color': FAINT})],
     [('Validation · financials · recommendations', {'size': 7.5, 'color': FAINT})]],
    align=PP_ALIGN.RIGHT, space_after=1.5)
box(s, M, Inches(0.80), prs.slide_width-2*M, Pt(2.2), fill=NAVY)

# ---- band A left ----
AX, AW = M, Inches(6.55)
y = sec(s, AX, Inches(0.94), AW, "The model, run on Jindal's disclosed reality")
tb = mktable(s, 7, 5, AX, y, AW, Inches(2.0))
tb.columns[0].width = Inches(1.95); tb.columns[1].width = Inches(1.0)
tb.columns[2].width = Inches(1.05); tb.columns[3].width = Inches(1.4); tb.columns[4].width = Inches(1.15)
data = [('Lever / result','FY24 (peak)','FY26 (actual)','FY28–30 (modelled)','Source'),
        ('Scrap charge','~60%','70.1%','78% (cap 80%)','BRSR FY26'),
        ('Renewable electricity','~30%','47%','65%','ESG Factsheet'),
        ('Green H₂ / Bio-LDO','–','pilot (Hisar)','scaled (Jajpur)','company disclosures'),
        ('Scope 1+2 intensity (tCO₂e/t)','2.15','1.76','1.35–1.45','ESG Factsheet / CDP'),
        ('Scope 3 intensity','1.90','1.27','0.95–1.05','ESG Factsheet'),
        ('Cradle-to-gate total','4.05','3.03','≈2.30–2.50','derived')]
for i,row in enumerate(data):
    kpi = i >= 4
    for j,v in enumerate(row):
        col = GRAY if i==0 else TEXT
        if kpi and j==2: col = BLUE
        if kpi and j==3: col = AMBERD
        if i>0 and j==4: col = FAINT
        set_cell(tb.cell(i,j), [[(v, {'size':7.2 if j==4 else 8,'bold':(i==0) or (kpi and j<4),
                 'color':col})]],
                 fill=THBG if i==0 else WHITE,
                 align=PP_ALIGN.LEFT if j in (0,4) else PP_ALIGN.CENTER)
for r in tb.rows: r.height = Inches(0.28)
y += Inches(2.24)
txt(s, AX, y, AW, Inches(0.3),
    [[('FY24→FY26 achieved: ', {'size': 8.8, 'color': TEXT}),
      ('−25.2% total footprint (−18.1% Scope 1+2)', {'size': 8.8, 'bold': True, 'color': GREEN}),
      ('; modelled path stays on the FY35 target glide of ', {'size': 8.8, 'color': TEXT}),
      ('0.99', {'size': 8.8, 'bold': True, 'color': NAVY}),
      (' (−50% vs FY22 baseline).', {'size': 8.8, 'color': TEXT})]], line_spacing=1.05)
y += Inches(0.36)
txt(s, AX, y, AW, Inches(0.3),
    [[('Every modelled lever is one Jindal is already investing in (₹700 Cr solar, 300 MW hybrid, hydrogen scale-up); the tool ranks them by tCO₂ per rupee of effort.', {'size': 7.8, 'color': FAINT, 'italic': True})]], line_spacing=1.0)
y += Inches(0.36)
vals = [('Validation 1 — audited anchor: ', 'model returns 2.8–3.1 for FY26 inputs; the audited 3.03 falls inside.'),
        ('Validation 2 — industry curve: ', 'results track the worldstainless scrap curve (85% → 1.95; 50% → 3.70; 30% → 6.80).'),
        ('Validation 3 — grade physics: ', 'nickel-free 430 lands near half the 304 footprint, as metallurgy requires.')]
for h,b in vals:
    txt(s, AX, y, AW, Inches(0.26),
        [[('▪ ', {'size': 8, 'color': GREEN}), (h, {'size': 8.4, 'bold': True, 'color': NAVY}),
          (b, {'size': 8.4, 'color': TEXT})]], line_spacing=1.0)
    y += Inches(0.265)

# ---- band A right: CBAM ----
BX = M + AW + Inches(0.35); BW = prs.slide_width - BX - M
y = sec(s, BX, Inches(0.94), BW, 'What verified data is worth at the EU border (2030)')
bars = [('EU default assigned to India (8.44 tCO₂/t incl. markup)', 446, RED, '€446/t'),
        ('India BF-BOF average', 122, STEEL, '€107–122'),
        ('Jindal, verified — conservative basis (1.4 t)', 74, BLUE, '€74'),
        ('Jindal, verified — CBAM-direct basis (0.9 t)', 48, GREEN, '€48')]
for lab, v, cc, vl in bars:
    txt(s, BX, y, BW, Inches(0.18), [[(lab, {'size': 8, 'color': TEXT, 'bold': cc in (BLUE,GREEN)})]])
    blen = Emu(int(Inches(3.6) * (v/446.0))) + Inches(0.05)
    box(s, BX, y+Inches(0.17), blen, Inches(0.13), fill=cc)
    txt(s, BX+blen+Inches(0.06), y+Inches(0.135), Inches(1.0), Inches(0.2),
        [[(vl, {'size': 8.5, 'bold': True, 'color': cc})]])
    y += Inches(0.405)
txt(s, BX, y, BW, Inches(0.2),
    [[('Mechanics: default 6.49 base + 30% no-data markup; phase-in 2.5%→48.5%; ETS €80→€109 (EU Regs 2025/2620-21).', {'size': 7.2, 'color': FAINT, 'italic': True})]])
y += Inches(0.26)
hline(s, BX, y, BW); y += Inches(0.06)
txt(s, BX, y, BW, Inches(0.42),
    [[('₹200–290 crore / year ', {'size': 21, 'bold': True, 'color': NAVY}),
      (' (€22–32M)', {'size': 11, 'bold': True, 'color': GRAY})]])
y += Inches(0.44)
txt(s, BX, y, BW, Inches(0.5),
    [[('€372–398 saved per tonne × 60–80 kt/yr of EU exports, earned by filing verified plant data instead of accepting EU defaults. The calculator produces exactly this verified figure per tonne.', {'size': 8.4, 'color': GRAY})]], line_spacing=1.05)

# ---- band B ----
bY = Inches(5.00)
y = sec(s, M, bY, AW, 'Initial recommendations')
recs = [('Operations: ', 'raise scrap 70→78% within the 80% tramp-element guardrail; deploy XRF/LIBS sensor sorting to secure clean feed.'),
        ('Operations: ', 'lift renewables 47→65% via the 300 MW hybrid; captive coal at 1.045 t/MWh is the second-largest lever.'),
        ('Finance: ', 'file verified MRV data with EU importers — a €372–398/t saving vs defaults; the tool generates the per-tonne statement.'),
        ('Sales: ', 'attach a carbon statement to every EU order; CPR GWP declarations are mandatory since Jan 2026 and the Digital Product Passport follows ~2028.')]
for i,(h,b) in enumerate(recs, start=1):
    txt(s, M, y, AW, Inches(0.36),
        [[(f'{i}.  ', {'size': 8.3, 'bold': True, 'color': BLUE}),
          (h, {'size': 8.3, 'bold': True, 'color': NAVY}), (b, {'size': 8.3, 'color': GRAY})]],
        line_spacing=1.0)
    y += Inches(0.355)
txt(s, M, y+Inches(0.01), AW, Inches(0.24),
    [[('Combined effect: Scope 1+2 falls to 1.35–1.45 by FY28-30 — on track for the FY35 −50% target.', {'size': 8.4, 'bold': True, 'color': GREEN})]])

y = sec(s, BX, bY, BW, 'Deployment roadmap')
tb = mktable(s, 5, 6, BX, y, BW, Inches(1.28))
tb.columns[0].width = Inches(2.55)
for j in range(1,6): tb.columns[j].width = Emu(int((BW-Inches(2.55))/5))
rm_rows = [('Calculator live; verified CBAM filing', GREEN, [1,2]),
           ('AI scrap sorting (raise the 80% ceiling)', BLUE, [2,3]),
           ('Biochar in ferrochrome; slag to cement', AMBERD, [3,4]),
           ('Product passport per heat (EU DPP)', BLUE, [4,5])]
for j,h in enumerate(['Initiative','FY26','FY27','FY28','FY29','FY30']):
    set_cell(tb.cell(0,j), [[(h, {'size':7.4,'bold':True,'color':GRAY})]], fill=THBG,
             align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
for i,(name, cc, cols_) in enumerate(rm_rows, start=1):
    set_cell(tb.cell(i,0), [[(name, {'size':7.6,'color':TEXT})]], fill=WHITE)
    for j in range(1,6):
        set_cell(tb.cell(i,j), '', fill=cc if j in cols_ else WHITE)
for r in tb.rows: r.height = Inches(0.22)
y += Inches(1.18)
y = sec(s, BX, y, BW, 'Assumptions & limitations')
txt(s, BX, y-Inches(0.03), BW, Inches(0.4),
    [[('Per tonne of crude steel; Scope 2 market-based; CBAM prices direct emissions only (conservative). FeMo ±30%; NPI 60–85 t/t; ETS path (€80→€109) is analyst consensus — savings scale with realised price.', {'size': 7.4, 'color': GRAY})]],
    line_spacing=1.05)

footer(s, 2, 'Source: JSL ESG Factsheets, BRSR, CDP FY22–26; EU Regulations 2023/956, 2025/2620-21, 2026/1740; CEA v21.0; worldstainless/ISSF; ICDA; Goldman Sachs, ICRA, CRISIL analyst reports; team analysis.')

out = os.path.join(HERE, 'StainlessCarbon_Round1.pptx')
prs.save(out)
print('saved', out)
